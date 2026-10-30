"""Read-only reader for the Ridhi keyword research workbook kept in data/.

The owner drops `ridhi-keyword-db-YYYY-MM-DD.xlsx` into the data folder; Flowmerce reads the newest such file and
shows its rows on the Topics & keywords page. Nothing is written to the database and the file is never modified.
Figures are exactly what the sheet says (Semrush free-plan estimates, Google Trends relative interest): numbers
are only extracted from text such as "1,300 (Semrush US 2026-10-19)", never invented. Ridhi only.
"""
from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..config import config

BRAND = "ridhi"
PATTERN = "ridhi-keyword-db-*.xlsx"
MAIN_SHEET = "Keyword DB"
NEW_SHEET = "New keywords"
OPP_SHEET = "Opportunities"

_cache: dict[str, Any] = {"key": None, "value": None}


def _num(v: Any) -> float | None:
    """The leading number of a cell ("1,300 (Semrush ...)" -> 1300, "$0.80; ..." -> 0.8, "8.12" -> 8.12)."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.match(r"\s*\$?\s*(-?[\d,]*\.?\d+)", str(v))
    if not m:
        return None
    try:
        return float(m.group(1).replace(",", ""))
    except ValueError:
        return None


def _int(v: Any) -> int | None:
    n = _num(v)
    return None if n is None else int(round(n))


def _text(v: Any) -> str | None:
    if v is None:
        return None
    t = str(v).strip()
    return t or None


def find_file() -> Path | None:
    """The newest matching workbook (the date is in the file name), or None."""
    folder = config.root / "data"
    files = sorted(p for p in folder.glob(PATTERN) if p.is_file() and not p.name.startswith("~$"))
    return files[-1] if files else None


def _rows(ws: Any) -> list[dict[str, Any]]:
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if not header:
        return []
    names = [str(h).strip() if h is not None else "" for h in header]
    out = []
    for row in it:
        if not any(c is not None for c in row):
            continue
        out.append({n: c for n, c in zip(names, row, strict=False) if n})
    return out


def _parse(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = set(wb.sheetnames)
        if MAIN_SHEET not in sheets:
            return {"error": f"The workbook has no “{MAIN_SHEET}” sheet.", "rows": []}
        main = _rows(wb[MAIN_SHEET])
        new_ids = {str(r.get("id")) for r in _rows(wb[NEW_SHEET])} if NEW_SHEET in sheets else set()
        opp_ids = {str(r.get("id")) for r in _rows(wb[OPP_SHEET])} if OPP_SHEET in sheets else set()
        title = None
        if "Summary" in sheets:
            first = next(wb["Summary"].iter_rows(values_only=True, max_row=1), None)
            title = _text(first[0]) if first else None
    finally:
        wb.close()
    rows = []
    for r in main:
        kid = _text(r.get("id"))
        kw = _text(r.get("keyword"))
        if not kid or not kw:
            continue
        rows.append({
            "id": kid,
            "keyword": kw,
            "topic": _text(r.get("topic")),
            "cluster": _text(r.get("cluster")),
            "intent": _text(r.get("intent")),
            "journey_stage": _text(r.get("journey_stage")),
            "search_volume": _int(r.get("search_volume")),
            "keyword_difficulty": _int(r.get("keyword_difficulty")),
            "cpc": _num(r.get("cpc_competition")),
            "gsc_clicks": _int(r.get("gsc_clicks")),
            "gsc_impressions": _int(r.get("gsc_impressions")),
            "gsc_avg_position": _num(r.get("gsc_avg_position")),
            "existing_page": _text(r.get("existing_ridhi_page")),
            "cannibalization_risk": _text(r.get("cannibalization_risk")),
            "recommended_action": _text(r.get("recommended_action")),
            "confidence": _text(r.get("confidence")),
            "notes": _text(r.get("notes")),
            "volume_source": _text(r.get("search_volume")),
            "new": kid in new_ids,
            "opportunity": kid in opp_ids,
        })
    return {"title": title, "rows": rows, "error": None}


def load() -> dict[str, Any]:
    """The parsed workbook, re-read only when the file changes."""
    path = find_file()
    if path is None:
        return {"found": False, "brand": BRAND, "pattern": PATTERN, "folder": "data/", "rows": []}
    st = path.stat()
    key = (str(path), st.st_mtime_ns, st.st_size)
    if _cache["key"] != key:
        try:
            parsed = _parse(path)
        except Exception as e:  # unreadable or open in Excel with a lock: report it, do not crash
            parsed = {"title": None, "rows": [], "error": f"Could not read {path.name}: {e}"}
        _cache["key"] = key
        _cache["value"] = parsed
    parsed = _cache["value"]
    clusters = sorted({r["cluster"] for r in parsed["rows"] if r["cluster"]})
    return {
        "found": True,
        "brand": BRAND,
        "file": path.name,
        "modifiedAt": datetime.fromtimestamp(st.st_mtime, tz=timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
        "title": parsed.get("title"),
        "error": parsed.get("error"),
        "clusters": clusters,
        "rows": parsed["rows"],
    }
